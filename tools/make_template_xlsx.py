#!/usr/bin/env python3
"""Tạo mau-du-lieu/Mau_nhap_lieu_KDI.xlsx (có danh sách chọn) từ cấu trúc bảng trong apps-script/.

Chạy:  python3 tools/make_template_xlsx.py   (cần: pip install openpyxl, Node >= 18)
"""
import json
import os
import subprocess

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
JS = """
const { createGasEnv } = require('./tests/gas-env');
const { ctx } = createGasEnv();
const spec = ctx.importSpec_({}, { role: 'admin', perms: [] , isSuperAdmin: true });
process.stdout.write(JSON.stringify({ spec, weekdays: ctx.WEEKDAYS }));
"""


def load_spec():
    data = json.loads(subprocess.check_output(["node", "-e", JS], cwd=ROOT))
    return data["spec"], data["weekdays"]


HEAD_FILL = PatternFill("solid", fgColor="1F3B73")
REQ_FILL = PatternFill("solid", fgColor="C2410C")
EX_FONT = Font(italic=True, color="667085")
MAX_ROWS = 3000
REF_SHEETS = {"school": "TruongHoc", "teacher": "GiaoVien", "assistant": "TroGiang", "subject": "MonHoc"}


def build(spec, weekdays, rows_by_table=None, out=None, intro=None):
    """rows_by_table: {table: [[giá trị theo thứ tự cột], ...]} – None = dùng dòng ví dụ."""
    wb = Workbook()
    guide = wb.active
    guide.title = "HuongDan"
    lines = (intro or []) + [
        "HƯỚNG DẪN NHẬP DỮ LIỆU – KDI SCHEDULER",
        "",
        "1. Mỗi sheet là 1 bảng dữ liệu. Giữ nguyên TÊN SHEET và DÒNG TIÊU ĐỀ (dòng 1). Cột tiêu đề màu cam là BẮT BUỘC.",
        "2. Các dòng chữ nghiêng màu xám là VÍ DỤ – hãy xóa hoặc ghi đè bằng dữ liệu thật.",
        "3. Sheet không cần nhập thì để trống (chỉ còn dòng tiêu đề).",
        "4. Thứ tự hệ thống nhập: " + " → ".join(t["sheet"] for t in spec) + ".",
        "5. Mã (môn, trường, GV) viết liền không dấu, vd. G01, TR001, GV001, TG001. Trùng mã đã có trong hệ thống = cập nhật bản ghi đó.",
        "6. Ô có danh sách thả xuống (Thứ, Loại GV, Xếp loại, Ưu tiên, Vai trò, Mã trường, Mã GV, Mã TG, Mã môn…): bấm vào ô để CHỌN, không gõ tay.",
        "   Danh sách mã lấy từ các sheet MonHoc, TruongHoc, GiaoVien, TroGiang trong file – thêm dòng ở đó là chọn được ở sheet khác.",
        "7. Vị trí trường / nhà GV / nhà TG: Google Maps → chọn địa điểm → Chia sẻ → Sao chép đường liên kết → dán vào cột 'Link Google Maps'.",
        "8. Lớp học chương trình gì: khai báo ở LopHoc. GV đã được đào tạo chương trình nào: GV_DaoTao (mỗi dòng 1 GV × 1 chương trình).",
        "9. Lưu file → vào web → 'Nhập dữ liệu từ Excel' → chọn file → Kiểm tra → Nhập dữ liệu.",
        "",
    ]
    for i, text in enumerate(lines, 1):
        guide.cell(row=i, column=1, value=text)
    guide["A1"].font = Font(bold=True, size=14, color="1F3B73")
    r = len(lines) + 1
    for c, h in enumerate(["Sheet", "Cột", "Bắt buộc", "Cách ghi"], 1):
        cell = guide.cell(row=r, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEAD_FILL
    for t in spec:
        for i, col in enumerate(t["columns"]):
            r += 1
            guide.cell(row=r, column=1, value=(t["sheet"] + " – " + t["label"]) if i == 0 else "")
            guide.cell(row=r, column=2, value=col["header"])
            guide.cell(row=r, column=3, value="Có" if col["required"] else "")
            guide.cell(row=r, column=4, value=col.get("hint") or "")
    for col, w in zip("ABCD", [36, 24, 10, 80]):
        guide.column_dimensions[col].width = w

    for t in spec:
        ws = wb.create_sheet(t["sheet"])
        for c, col in enumerate(t["columns"], 1):
            cell = ws.cell(row=1, column=c, value=col["header"])
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = REQ_FILL if col["required"] else HEAD_FILL
            cell.alignment = Alignment(vertical="center")
            letter = get_column_letter(c)
            ws.column_dimensions[letter].width = max(12, min(40, len(col["header"]) + 6))
            options = col.get("options")
            if col["type"] == "weekday":
                options = weekdays
            elif col["type"] == "bool":
                options = ["TRUE", "FALSE"]
            formula = None
            if options:
                formula = '"' + ",".join(options) + '"'
            elif col["type"] in REF_SHEETS and REF_SHEETS[col["type"]] != t["sheet"]:
                # Danh sách thả xuống lấy từ cột mã ở sheet khác trong cùng file.
                formula = "=" + REF_SHEETS[col["type"]] + "!$A$2:$A$" + str(MAX_ROWS)
            if formula:
                dv = DataValidation(type="list", formula1=formula, allow_blank=True)
                dv.error = "Hãy chọn giá trị trong danh sách"
                dv.errorTitle = "Giá trị không hợp lệ"
                dv.showErrorMessage = True
                ws.add_data_validation(dv)
                dv.add(f"{letter}2:{letter}{MAX_ROWS}")
            if col["type"] in ("time", "phone", "weekday") or col["key"] in ("code", "teacher", "school", "subject", "id"):
                for row in range(2, MAX_ROWS + 1):
                    ws.cell(row=row, column=c).number_format = "@"
        sample = rows_by_table is None
        rows = t["examples"] if sample else rows_by_table.get(t["table"], [])
        for r_i, ex in enumerate(rows, 2):
            for c, v in enumerate(ex, 1):
                cell = ws.cell(row=r_i, column=c, value=v)
                if sample:
                    cell.font = EX_FONT
        ws.freeze_panes = "A2"

    out = out or os.path.join(ROOT, "mau-du-lieu", "Mau_nhap_lieu_KDI.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    spec, weekdays = load_spec()
    print("Đã tạo", build(spec, weekdays))
